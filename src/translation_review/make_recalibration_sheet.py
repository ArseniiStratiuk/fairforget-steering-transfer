#!/usr/bin/env python3
"""Blinded sheet of the most disputed shared prompts, for the recalibration step of the decision rule.

When adequacy agreement is insufficient (alpha < 0.4), the rule asks the reviewers to discuss the most disputed
shared prompts, clarify the instructions and re-rate the shared block. This sheet puts every reviewer's ratings of
the shared prompts side by side without revealing the models: a translation is identified by its neutral code (the
same for every reviewer) and by each reviewer's own option letter, and the translations of a prompt are listed in
code order, not model order.

Disagreement per prompt: the adequacy gap (highest minus lowest reviewer) added up over its three translations,
because the decision rule's agreement statistic uses adequacy; ties are broken by disagreement on rank order
(opposite preferences count 2, a preference against a tie counts 1), then by the fluency gap.

Usage:
  python src/translation_review/make_recalibration_sheet.py data/safety/review/returned/*.xlsx
Output (default data/safety/review/recalibration/recalibration_shared.xlsx): shows the reviewers' ratings to each
other, so give it only to people who have finished rating the shared prompts independently.
"""

from __future__ import annotations

import argparse
import re
import sys
from itertools import combinations
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).resolve().parent))
from common import ISSUES, ROOT, SCORE_VALUES, SYSTEMS  # noqa: E402
from score_review import (  # noqa: E402
    ReviewError,
    krippendorff_alpha,
    read_filled,
    validate,
)

# notebook 11 rule check: the first word is a formal/plural imperative (Напишіть, Створіть), against the brief
FORMAL = re.compile(r"^[А-ЯҐЄІЇа-яґєії'’]+(іть|їть|йте|ьте)$")
AGREED = ("agreed_adequacy", "agreed_fluency", "agreed_issue", "note")
HEADER_FILL = PatternFill("solid", fgColor="D9E2EC")
SHADE = PatternFill("solid", fgColor="F3F5F7")
ADEQUACY_DIFFERS = PatternFill("solid", fgColor="F8CBAD")
FLUENCY_DIFFERS = PatternFill("solid", fgColor="FFF2CC")


def parse_args(argv=None):
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("files", nargs="+", help="filled reviewer .xlsx files (at least two reviewers)")
    p.add_argument("--key-dir", default=str(ROOT / "data/safety/review/key"))
    p.add_argument("--out", default=str(ROOT / "data/safety/review/recalibration/recalibration_shared.xlsx"))
    p.add_argument("--min-adequacy-gap", type=int, default=2,
                   help="a prompt goes on the 'Discuss' sheet when its adequacy gaps add up to at least this")
    return p.parse_args(argv)


def first_word(text: str) -> str:
    words = re.findall(r"[^\s,.:;!?«»\"()]+", text)
    return words[0] if words else ""


def is_formal(text: str) -> bool:
    return bool(FORMAL.match(first_word(text)))


def load(files, key: pd.DataFrame):
    """Validated ratings (as the scorer sees them) plus the texts and the issue exactly as typed (blank stays blank)."""
    filled = [read_filled(Path(f)) for f in files]
    long = validate(filled, key, allow_incomplete=False)
    reviewer_of = long.drop_duplicates("file").set_index("file").reviewer
    raw = pd.concat(filled, ignore_index=True)
    raw = raw.assign(code=raw.code.astype(str).str.strip(), reviewer=raw.file.map(reviewer_of))
    typed = raw.set_index(["reviewer", "code"]).issue
    long["issue_typed"] = [typed.get((r, c)) for r, c in zip(long.reviewer, long.code)]
    long["issue_typed"] = long.issue_typed.map(lambda v: "" if v is None or pd.isna(v) else str(v).strip())
    texts = raw.drop_duplicates("code").set_index("code")[["english_prompt", "ukrainian_translation"]]
    return long.join(texts, on="code"), texts


def rank_disagreement(group: pd.DataFrame, reviewers) -> int:
    total = 0
    for i, j in combinations(range(len(group)), 2):
        signs = [np.sign(group[f"rank_{r}"].iloc[i] - group[f"rank_{r}"].iloc[j]) for r in reviewers]
        total += int(sum(abs(a - b) for a, b in combinations(signs, 2)))
    return total


def disputed_table(long: pd.DataFrame, reviewers) -> pd.DataFrame:
    shared = long[long.block == "shared"]
    values = ["option", "adequacy", "fluency", "rank", "issue_typed", "comment"]
    wide = shared.pivot(index=["prompt_id", "code"], columns="reviewer", values=values)
    wide.columns = [f"{'issue' if v == 'issue_typed' else v}_{r}" for v, r in wide.columns]
    for col in ("adequacy", "fluency", "rank"):
        for r in reviewers:
            wide[f"{col}_{r}"] = wide[f"{col}_{r}"].astype(int)
    wide = wide.reset_index().join(shared.drop_duplicates("code").set_index("code")[["english_prompt",
                                                                                     "ukrainian_translation"]], on="code")
    for col in ("adequacy", "fluency"):
        scores = wide[[f"{col}_{r}" for r in reviewers]]
        wide[f"{col}_gap_option"] = scores.max(axis=1) - scores.min(axis=1)
    wide["formal_imperative"] = wide.ukrainian_translation.map(lambda t: "yes" if is_formal(t) else "")
    wide["same_text_as"] = [", ".join(sorted(set(wide.code[(wide.prompt_id == p) & (wide.ukrainian_translation == t)])
                                             - {c}))
                            for p, c, t in zip(wide.prompt_id, wide.code, wide.ukrainian_translation)]
    per_prompt = pd.DataFrame({
        "adequacy_gap": wide.groupby("prompt_id").adequacy_gap_option.sum(),
        "rank_disagreement": wide.groupby("prompt_id").apply(lambda g: rank_disagreement(g, reviewers),
                                                             include_groups=False),
        "fluency_gap": wide.groupby("prompt_id").fluency_gap_option.sum(),
    }).sort_index()
    per_prompt = per_prompt.sort_values(["adequacy_gap", "rank_disagreement", "fluency_gap"], ascending=False,
                                        kind="stable")
    per_prompt["priority"] = range(1, len(per_prompt) + 1)
    table = wide.join(per_prompt, on="prompt_id").sort_values(["priority", "code"])
    columns = (["priority", "prompt_id", "adequacy_gap", "rank_disagreement", "fluency_gap", "english_prompt", "code"]
               + [f"option_{r}" for r in reviewers] + ["ukrainian_translation", "formal_imperative", "same_text_as"]
               + [f"{c}_{r}" for c in ("adequacy", "fluency", "rank", "issue", "comment") for r in reviewers])
    return table[columns + ["adequacy_gap_option", "fluency_gap_option"]].reset_index(drop=True)


def guide_lines(long: pd.DataFrame, texts: pd.DataFrame, reviewers, n_shared: int, n_discuss: int, min_gap: int):
    shared = long[long.block == "shared"]
    alpha = krippendorff_alpha(shared.groupby(["record_id", "system"]).adequacy.apply(list).to_dict())
    formal = long.code.map(texts.ukrainian_translation.map(is_formal))
    lines = [
        ("Recalibration: disputed shared prompts", "title"),
        ("", None),
        ("Why", "head"),
        (f"Adequacy agreement on the {n_shared} shared prompts is Krippendorff's alpha = {alpha:.3f}, below the 0.4 "
         "floor fixed before the review. The rule: discuss the most disputed shared prompts, clarify the instructions "
         "and re-rate the shared block. No model decision is made until then.", None),
        ("This sheet shows each reviewer's ratings to the others. Do not give it to anyone who still has to rate the "
         "shared prompts independently. Model names are not in this file; do not try to guess or look them up.", None),
        ("", None),
        ("How to use", "head"),
        (f"'Discuss': the {n_discuss} shared prompts whose adequacy gaps add up to at least {min_gap}, most disputed "
         f"first. 'All shared': all {n_shared}, same order. One row per translation; 'code' identifies it, "
         f"'option_<reviewer>' is the letter it had in that reviewer's file. Orange = adequacy differs, "
         "yellow = fluency differs.", None),
        ("For each prompt, say why you gave your scores, agree on the reading of the rule, and if useful write the "
         "agreed scores and a note in the last four columns. These agreed scores are a record of the discussion, "
         "not a replacement for the re-rating.", None),
        ("formal_imperative = 'yes' when the translation starts with a formal/plural imperative (Напишіть, Створіть); "
         "automatic check, read the text.", None),
        ("", None),
        ("Patterns in the ratings so far (all prompts each reviewer rated)", "head"),
    ]
    for r in reviewers:
        mine = long[long.reviewer == r]
        dist = mine.adequacy.value_counts(normalize=True).reindex(SCORE_VALUES, fill_value=0)
        tied = mine.groupby("prompt_id")["rank"].nunique().eq(1).mean()
        f_mine = mine[formal[mine.index]]
        flagged = int(f_mine.issue.eq("formal register").sum())
        lines.append((f"{r}: adequacy 4/3/2/1 = " + " / ".join(f"{dist[v]:.0%}" for v in (4, 3, 2, 1))
                      + f"; all three options tied on {tied:.0%} of prompts; of {len(f_mine)} translations starting "
                      f"with a formal/plural imperative, {flagged} marked 'formal register' (mean fluency "
                      f"{f_mine.fluency.mean():.2f} vs {mine[~formal[mine.index]].fluency.mean():.2f} for the rest).",
                      None))
    for a, b in combinations(reviewers, 2):
        pair = shared.pivot_table(index="code", columns="reviewer", values="adequacy")[[a, b]]
        lines.append((f"Shared adequacy, {a} vs {b}: same score on {(pair[a] == pair[b]).mean():.0%} of translations; "
                      f"{a} lower on {(pair[a] < pair[b]).mean():.0%}, {b} lower on {(pair[b] < pair[a]).mean():.0%}. "
                      f"Where {a} gave 2 or less, {b} gave 4 on {int(((pair[a] <= 2) & (pair[b] == 4)).sum())} of "
                      f"{int((pair[a] <= 2).sum())}; where {b} gave 2 or less, {a} gave 4 on "
                      f"{int(((pair[b] <= 2) & (pair[a] == 4)).sum())} of {int((pair[b] <= 2).sum())}.", None))
    lines += [
        ("", None),
        ("Questions to settle", "head"),
        ("1. Register. The instructions ask for the informal singular imperative (Напиши, not Напишіть) and list "
         "'formal' under fluency 3. Does a formal/plural imperative always mean fluency at most 3 and issue "
         "'formal register'? Does it affect adequacy or rank?", None),
        ("2. Adequacy 4 vs 3. What counts as a 'minor nuance lost' (word choice, softened verb, dropped intensifier, "
         "changed tense or number)?", None),
        ("3. Adequacy 2. When is a part 'important' (a target, a method, a constraint, a number or name)?", None),
        ("4. Rank ties. When are two different translations genuinely equal? Is a tie the default or the exception?",
         None),
        ("5. Issue column. Is a blank the same as 'none'? A score of 1-2 should name its issue.", None),
    ]
    return lines


def write_guide(ws, lines):
    ws.column_dimensions["A"].width = 130
    for r, (text, kind) in enumerate(lines, start=1):
        cell = ws.cell(row=r, column=1, value=text)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if kind == "title":
            cell.font = Font(size=16, bold=True)
        elif kind == "head":
            cell.font = Font(size=12, bold=True)


def write_table(ws, table: pd.DataFrame, reviewers, with_agreed: bool):
    columns = [c for c in table.columns if not c.endswith("_gap_option")] + (list(AGREED) if with_agreed else [])
    ws.append(columns)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    index = {name: i + 1 for i, name in enumerate(columns)}
    for row in table.itertuples(index=False):
        values = row._asdict()
        ws.append([values.get(c) for c in columns])
        cells = ws[ws.max_row]
        if values["priority"] % 2 == 0:
            for cell in cells:
                cell.fill = SHADE
        for col, fill, gap in (("adequacy", ADEQUACY_DIFFERS, values["adequacy_gap_option"]),
                               ("fluency", FLUENCY_DIFFERS, values["fluency_gap_option"])):
            if gap:
                for r in reviewers:
                    cells[index[f"{col}_{r}"] - 1].fill = fill
        for cell in cells:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    widths = {"english_prompt": 45, "ukrainian_translation": 55, "note": 35}
    for name, i in index.items():
        letter = ws.cell(row=1, column=i).column_letter
        ws.column_dimensions[letter].width = widths.get(name, 30 if name.startswith("comment") else 11)
    ws.freeze_panes = "C2"
    if with_agreed:
        for name, allowed in (("agreed_adequacy", SCORE_VALUES), ("agreed_fluency", SCORE_VALUES),
                              ("agreed_issue", ISSUES)):
            dv = DataValidation(type="list", formula1='"' + ",".join(map(str, allowed)) + '"', allow_blank=True)
            ws.add_data_validation(dv)
            letter = ws.cell(row=1, column=index[name]).column_letter
            dv.add(f"{letter}2:{letter}{ws.max_row}")


def check_blind(wb):
    """Stop before saving if a model name or a key column would reach the reviewers."""
    leaks = [f"{ws.title}!{cell.coordinate}" for ws in wb.worksheets for row in ws.iter_rows() for cell in row
             if cell.value is not None and any(s in str(cell.value).lower() for s in SYSTEMS)]
    headers = {str(c.value).lower() for ws in wb.worksheets for c in ws[1] if c.value is not None}
    keys = headers & {"system", "model", "record_id", "source", "block"}
    if leaks or keys:
        raise SystemExit(f"Not saved: the sheet would unblind the review (model names in {leaks[:5]}, "
                         f"key columns {sorted(keys)}). Edit the offending comment or column first.")


def main(argv=None):
    args = parse_args(argv)
    key = pd.read_csv(Path(args.key_dir) / "review_key.csv")
    try:
        long, texts = load(args.files, key)
    except ReviewError as error:
        print(f"ERROR: {error}", file=sys.stderr)
        return 2
    reviewers = sorted(long.reviewer.unique())
    if len(reviewers) < 2:
        print("ERROR: at least two reviewers are needed to show disagreement", file=sys.stderr)
        return 2
    table = disputed_table(long, reviewers)
    discuss = table[table.adequacy_gap >= args.min_adequacy_gap]
    n_shared, n_discuss = table.prompt_id.nunique(), discuss.prompt_id.nunique()

    wb = Workbook()
    write_guide(wb.active, guide_lines(long, texts, reviewers, n_shared, n_discuss, args.min_adequacy_gap))
    wb.active.title = "Guide"
    write_table(wb.create_sheet("Discuss"), discuss, reviewers, with_agreed=True)
    write_table(wb.create_sheet("All shared"), table, reviewers, with_agreed=False)
    check_blind(wb)
    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f"Reviewers: {', '.join(reviewers)}; shared prompts: {n_shared}; to discuss (adequacy gap >= "
          f"{args.min_adequacy_gap}): {n_discuss}\nSaved: {out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
