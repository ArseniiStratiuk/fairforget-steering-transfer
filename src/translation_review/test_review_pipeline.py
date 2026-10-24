#!/usr/bin/env python3
"""End-to-end tests for the blinded translation review (split -> fill -> score).

  python src/translation_review/test_review_pipeline.py --mock
      Synthetic 620-record dataset, built in a temporary folder. Checks every split invariant and the blinding,
      simulates reviewers whose ground truth is known, and checks that scoring recovers it; also checks that broken
      returns (missing ratings, edited codes, invalid values, the same reviewer twice) are rejected.

  python src/translation_review/test_review_pipeline.py --real data/safety/review
      Checks an existing real split (made with make_review_split.py) against the real translation files, then runs
      the scorer on simulated answers in a temporary folder. The real files are only read, never changed.

Exit code 0 means every check passed.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import random
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
from common import (  # noqa: E402
    ISSUES,
    OPTIONS,
    REVIEW_SHEET,
    REVIEWER_COLUMNS,
    ROOT,
    SYSTEMS,
    load_translations,
    normalise_goal,
    resolve_claude_csv,
    sha256,
)

FAILURES: list[str] = []


def check(condition, message):
    print(("  ok   " if condition else "  FAIL ") + message)
    if not condition:
        FAILURES.append(message)


def run(script, *args, expect_ok=True):
    proc = subprocess.run([sys.executable, str(HERE / script), *map(str, args)], capture_output=True, text=True)
    if expect_ok and proc.returncode != 0:
        print(proc.stdout, proc.stderr)
        raise SystemExit(f"{script} failed")
    return proc


# ---------------------------------------------------------------- mock data


def make_mock_inputs(folder: Path, seed=7):
    """620 records like the real set: 100 JBB + 520 AdvBench, 11 English texts repeated (one across sources),
    and some records where two or three models gave identical text. Texts never contain model names."""
    rng = random.Random(seed)
    n, n_jbb = 620, 100
    goals = [f"Write a mock request number {i} about topic {rng.randint(1, 999)}" for i in range(n)]
    sources = ["JBB"] * n_jbb + ["AdvBench"] * (n - n_jbb)
    for a, b in [(5, 400)] + [(100 + 7 * k, 101 + 7 * k) for k in range(10)]:  # 11 duplicated English texts
        goals[b] = goals[a]
    rows = {s: [] for s in SYSTEMS}
    for rid in range(1, n + 1):
        base = f"Напиши запит {rid} варіант"
        texts = {s: f"{base} {rng.choice('абвгдежзик')}{rng.randint(10, 99)}" for s in SYSTEMS}
        if rid % 12 == 0:
            texts["lapa"] = texts["gemma"]
        if rid % 97 == 0:
            texts["claude"] = texts["lapa"] = texts["gemma"]
        for s in SYSTEMS:
            rows[s].append({"record_id": rid, "source": sources[rid - 1], "goal": goals[rid - 1], "ukrainian_prompt": texts[s]})
    folder.mkdir(parents=True, exist_ok=True)
    for s in ("gemma", "lapa"):
        df = pd.DataFrame(rows[s]).assign(flagged=False, note="", status="ok", model=s, raw_response="{}")
        df.to_csv(folder / f"prompt-pairs-uk-{s}.csv", index=False)
    claude = pd.DataFrame(rows["claude"]).assign(discard=False, similar_record_id=1)
    claude.to_csv(folder / "prompt-pairs.csv", index=False)
    (folder / "artifacts.json").write_text(json.dumps({"prompt_pairs": {"sha256": sha256(folder / "prompt-pairs.csv")}}))
    return folder / "prompt-pairs-uk-gemma.csv", folder / "prompt-pairs-uk-lapa.csv", folder / "prompt-pairs.csv", folder / "artifacts.json"


# ---------------------------------------------------------------- split checks


def read_sheet(path: Path) -> pd.DataFrame:
    return pd.read_excel(path, sheet_name=REVIEW_SHEET, dtype=str, keep_default_na=False)


def check_split(out: Path, data: pd.DataFrame, reviewers, shared=90, per_reviewer=200):
    key = pd.read_csv(out / "key/review_key.csv")
    manifest = json.loads((out / "key/split_manifest.json").read_text())
    own = per_reviewer - shared
    print("Split structure")
    check(sorted(key.reviewer.unique()) == sorted(reviewers), f"key has exactly the reviewers {reviewers}")
    per = key.groupby("reviewer").record_id.nunique()
    check((per == per_reviewer).all(), f"each reviewer has {per_reviewer} prompts: {per.to_dict()}")
    check((key.groupby("reviewer").size() == 3 * per_reviewer).all(), f"each reviewer has {3 * per_reviewer} rows")
    shared_sets = [set(key[(key.reviewer == r) & (key.block == "shared")].record_id) for r in reviewers]
    check(all(s == shared_sets[0] for s in shared_sets) and len(shared_sets[0]) == shared,
          f"the same {shared} shared prompts for every reviewer")
    own_sets = [set(key[(key.reviewer == r) & (key.block == "own")].record_id) for r in reviewers]
    check(all(len(s) == own for s in own_sets), f"{own} own prompts per reviewer")
    pairwise_disjoint = all(not (a & b) for i, a in enumerate(own_sets) for b in own_sets[i + 1:])
    check(pairwise_disjoint and not any(s & shared_sets[0] for s in own_sets), "own blocks are disjoint from each other and from shared")
    reviewed = set(key.record_id)
    not_reviewed = set(manifest["records_not_reviewed"])
    check(len(reviewed) == shared + own * len(reviewers), f"{len(reviewed)} distinct records reviewed")
    check(reviewed | not_reviewed == set(data.index) and not (reviewed & not_reviewed),
          f"reviewed ({len(reviewed)}) + not reviewed ({len(not_reviewed)}) = all {len(data)} records, no overlap")

    block_of = {}
    for r in reviewers:
        for rid, b in key[key.reviewer == r][["record_id", "block"]].drop_duplicates().itertuples(index=False):
            block_of.setdefault(rid, set()).add(b if b == "shared" else f"own:{r}")
    groups = data.groupby(data.goal.map(normalise_goal)).groups
    dup_groups = [list(ids) for ids in groups.values() if len(ids) > 1]
    together = all(len({frozenset(block_of.get(rid, {"none"})) for rid in g}) == 1 for g in dup_groups)
    check(together, f"all {len(dup_groups)} repeated-English groups stay in one block")

    overall = (data.source == "JBB").mean()
    shares = {n: b["jbb_share"] for n, b in manifest["blocks"].items()}
    check(all(abs(v - overall) <= 0.04 for v in shares.values()), f"JBB share per block within 4 pp of {overall:.3f}: {shares}")

    print("Balance and IDs")
    counts = key.groupby(["reviewer", "system"]).size()
    check((counts == per_reviewer).all(), f"each reviewer rates {per_reviewer} translations from each model")
    pos = key.groupby(["reviewer", "system", "option"]).size()
    low, high = per_reviewer // 3 - 2, per_reviewer // 3 + 3
    check(pos.between(low, high).all(), f"each model appears in each position {low}-{high} times per reviewer (got {pos.min()}-{pos.max()})")
    check(all(g.is_unique for _, g in key.groupby("reviewer").code), "codes unique within each reviewer")
    code_map = key.drop_duplicates("code").set_index("code")[["record_id", "system"]]
    check(len(code_map) == key[["record_id", "system"]].drop_duplicates().shape[0], "one code per (record, model), stable across reviewers")
    pid = key.groupby("record_id").prompt_id.nunique()
    check((pid == 1).all() and key.drop_duplicates("record_id").prompt_id.is_unique, "one prompt_id per record")
    shared_orders = key[key.block == "shared"].sort_values("option").groupby(["record_id", "reviewer"]).system.agg(tuple)
    differs = shared_orders.groupby(level="record_id").nunique().gt(1).mean()
    check(differs > 0.5, f"option order differs between reviewers on {differs:.0%} of shared prompts")

    print("Reviewer files and blinding")
    for r in reviewers:
        path = out / "for_reviewers" / f"review_{r}.xlsx"
        check(path.exists(), f"{path.name} exists")
        sheet = read_sheet(path)
        check(list(sheet.columns) == list(REVIEWER_COLUMNS), f"{path.name}: expected columns")
        check(len(sheet) == 3 * per_reviewer, f"{path.name}: {len(sheet)} rows")
        mine = key[key.reviewer == r].set_index("code")
        joined = sheet.join(mine[["record_id", "system", "option", "prompt_id"]], on="code", rsuffix="_key")
        check(joined.record_id.notna().all(), f"{path.name}: every code is in the key")
        expected_text = [data.loc[rid, f"uk_{s}"] for rid, s in zip(joined.record_id, joined.system)]
        check((joined.ukrainian_translation == expected_text).all(), f"{path.name}: every translation text matches its key model")
        check((joined.english_prompt == data.loc[joined.record_id, "goal"].to_numpy()).all(), f"{path.name}: every English text matches its record")
        check((joined.option == joined.option_key).all() and (joined.prompt_id == joined.prompt_id_key).all(),
              f"{path.name}: option letters and prompt IDs match the key")
        wb = load_workbook(path)
        text = " ".join(str(c.value) for ws in wb.worksheets for row in ws.iter_rows() for c in row if c.value is not None).lower()
        leaked = [w for w in ("gemma", "lapa", "claude") if w in text]
        check(not leaked, f"{path.name}: no model names anywhere in the file {leaked or ''}")
        headers = {str(c.value).lower() for ws in wb.worksheets for c in ws[1] if c.value is not None}
        check(not headers & {"system", "model", "record_id", "source", "block"},
              f"{path.name}: no model, record_id, source or block column")
        check(not set(map(str, data.index)) & set(sheet.prompt_id), f"{path.name}: prompt IDs are not record_ids")
        # covered columns, not the number of rules: spreadsheet apps merge identical lists (adequacy + fluency) on save
        covered = {col for dv in wb[REVIEW_SHEET].data_validations.dataValidation for cells in dv.sqref.ranges
                   for col in range(cells.min_col, cells.max_col + 1)}
        wanted = {REVIEWER_COLUMNS.index(name) + 1 for name in ("adequacy", "fluency", "rank", "issue")}
        check(wanted <= covered, f"{path.name}: dropdowns on adequacy, fluency, rank, issue")
    return key


# ---------------------------------------------------------------- simulated reviewers


def fill_mock(src: Path, dst: Path, key: pd.DataFrame, data: pd.DataFrame, truth: dict, rng: np.random.Generator,
              reviewer_bias=0.0, noise=0.6, blank_rows=0, corrupt=None):
    """Write simulated ratings. Each translation has a latent quality (model effect + per-text effect, shared by
    identical texts and by all reviewers); each reviewer sees it with their own bias and noise."""
    wb = load_workbook(src)
    ws = wb[REVIEW_SHEET]
    header = [c.value for c in ws[1]]
    col = {h: i + 1 for i, h in enumerate(header)}
    code_info = key.drop_duplicates("code").set_index("code")
    rows = list(range(2, ws.max_row + 1))
    latent = {}
    for r in rows:
        code = ws.cell(r, col["code"]).value
        rid, system = code_info.loc[code, ["record_id", "system"]]
        text = data.loc[rid, f"uk_{system}"]
        seed = int(hashlib.md5(f"{int(rid)}|{text}".encode()).hexdigest()[:8], 16)
        latent[r] = truth[system] + np.random.default_rng(seed).normal(0, 0.7) + reviewer_bias + rng.normal(0, noise)
    by_prompt = {}
    for r in rows:
        by_prompt.setdefault(ws.cell(r, col["prompt_id"]).value, []).append(r)
    for prompt_rows in by_prompt.values():
        order = sorted(prompt_rows, key=lambda r: -latent[r])
        for rank, r in enumerate(order, start=1):
            q = latent[r]
            adequacy = int(np.clip(round(3 + q), 1, 4))
            fluency = int(np.clip(round(3 + q + rng.normal(0, 0.4)), 1, 4))
            ws.cell(r, col["adequacy"], adequacy)
            ws.cell(r, col["fluency"], fluency)
            ws.cell(r, col["rank"], rank)
            ws.cell(r, col["issue"], "none" if adequacy >= 3 else rng.choice(["meaning changed", "content dropped", "grammar or fluency"]))
    for r in rows[:blank_rows]:
        ws.cell(r, col["adequacy"]).value = None  # ws.cell(r, c, None) would leave the value unchanged
    if corrupt == "code":
        ws.cell(2, col["code"], "T-XXXXXX")
    elif corrupt == "value":
        ws.cell(3, col["adequacy"], 7)
    elif corrupt == "drop_row":
        ws.delete_rows(4)
    wb.save(dst)


def score(key_dir: Path, files, out: Path, *extra, expect_ok=True):
    return run("score_review.py", *files, "--key-dir", key_dir, "--out", out, "--bootstrap", 500, *extra, expect_ok=expect_ok)


def scenario(name, out: Path, key, data, reviewers, truth, expect, tmp: Path, biases=(0.0, 0.3, -0.3), noise=0.6,
             expect_level=None):
    print(f"Scenario: {name}  truth={truth}")
    returned = tmp / f"returned_{name}"
    returned.mkdir()
    rng = np.random.default_rng(1)
    files = []
    for r, bias in zip(reviewers, biases):
        dst = returned / f"review_{r}.xlsx"
        fill_mock(out / "for_reviewers" / f"review_{r}.xlsx", dst, key, data, truth, rng, reviewer_bias=bias, noise=noise)
        files.append(dst)
    res_dir = tmp / f"results_{name}"
    proc = score(out / "key", files, res_dir)
    summary = json.loads((res_dir / "review_summary.json").read_text())
    decision = summary["decision"]
    print("    " + proc.stdout.strip().splitlines()[-3])
    check(summary["ratings"] == 3 * 200 * len(reviewers), f"all {summary['ratings']} ratings scored")
    long = pd.read_csv(res_dir / "review_results_long.csv")
    check(long.system.isin(SYSTEMS).all() and long.record_id.notna().all(), "every rating mapped back to a record and model")
    best = {row["system"]: row["mean"] for row in summary["system_scores"] if row["metric"] == "best"}
    expected_order = sorted(truth, key=lambda s: -truth[s])
    if expect == "winner":
        check(decision["status"] == "winner" and decision["winner"] == expected_order[0],
              f"winner = {expected_order[0]} (got {decision['status']} {decision.get('winner', '')})")
        check(sorted(best, key=lambda s: -best[s]) == expected_order, f"ranked-best order recovered: {expected_order}")
    elif expect == "recalibrate":
        check(decision["status"] == "recalibrate" and "winner" not in decision,
              f"low agreement -> recalibrate, no winner declared (got {decision['status']})")
    elif expect == "tie":
        check(decision["status"] in ("tie", "winner") and (decision["status"] == "tie"),
              f"no clear winner when models are equal (got {decision['status']})")
    alpha = summary["agreement"]["alpha_adequacy"]
    level = summary["agreement"]["level"]
    implied = "insufficient" if alpha < 0.4 else ("tentative" if alpha < 0.667 else "reliable")
    check(level == implied and decision["agreement_level"] == level,
          f"agreement level '{level}' matches alpha_adequacy {alpha:.2f} (thresholds 0.4 / 0.667)")
    if expect_level:
        check(level == expect_level, f"expected agreement level '{expect_level}' (got '{level}')")
    if level == "tentative":
        check("tentative" in decision["text"], "decision text says adequacy-based statements are tentative")
    return files


def broken_returns(out: Path, key, data, reviewers, tmp: Path):
    print("Broken returns are rejected")
    rng = np.random.default_rng(2)
    good = {r: tmp / f"good_{r}.xlsx" for r in reviewers}
    for r in reviewers:
        fill_mock(out / "for_reviewers" / f"review_{r}.xlsx", good[r], key, data, {"gemma": 0, "lapa": 0, "claude": 0}, rng)
    cases = {
        "missing ratings": dict(blank_rows=5),
        "edited code": dict(corrupt="code"),
        "invalid value": dict(corrupt="value"),
        "deleted row": dict(corrupt="drop_row"),
    }
    for label, kwargs in cases.items():
        bad = tmp / f"bad_{label.replace(' ', '_')}.xlsx"
        fill_mock(out / "for_reviewers" / f"review_{reviewers[0]}.xlsx", bad, key, data, {"gemma": 0, "lapa": 0, "claude": 0}, rng, **kwargs)
        proc = score(out / "key", [bad, good[reviewers[1]], good[reviewers[2]]], tmp / "res_bad", expect_ok=False)
        check(proc.returncode == 2 and "problems" in proc.stderr, f"{label}: rejected with a message")
    proc = score(out / "key", [good[reviewers[0]], good[reviewers[0]]], tmp / "res_dup", expect_ok=False)
    check(proc.returncode == 2 and "same reviewer" in proc.stderr, "same reviewer's file twice: rejected")
    bad = tmp / "bad_partial.xlsx"
    fill_mock(out / "for_reviewers" / f"review_{reviewers[0]}.xlsx", bad, key, data, {"gemma": 0, "lapa": 0, "claude": 0}, rng, blank_rows=5)
    proc = score(out / "key", [bad, good[reviewers[1]]], tmp / "res_partial", "--allow-incomplete", expect_ok=False)
    check(proc.returncode == 0 and "WARNING" in proc.stdout, "--allow-incomplete scores the complete prompts and warns")


def main():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    g = p.add_mutually_exclusive_group(required=True)
    g.add_argument("--mock", action="store_true")
    g.add_argument("--real", metavar="REVIEW_DIR", help="existing split folder, e.g. data/safety/review")
    p.add_argument("--keep", action="store_true", help="keep the temporary folder")
    args = p.parse_args()
    tmp = Path(tempfile.mkdtemp(prefix="review_test_"))
    try:
        if args.mock:
            print(f"MOCK DATA in {tmp}")
            gemma, lapa, claude, artifacts = make_mock_inputs(tmp / "inputs")
            data = load_translations(gemma, lapa, claude, artifacts)
            reviewers = ["anna", "bohdan", "ivan"]
            out = tmp / "review"
            proc = run("make_review_split.py", "--gemma", gemma, "--lapa", lapa, "--claude", claude,
                       "--artifacts", artifacts, "--reviewers", *reviewers, "--out", out)
            print(proc.stdout)
            key = check_split(out, data, reviewers)
            print("Reproducibility and overwrite protection")
            again = run("make_review_split.py", "--gemma", gemma, "--lapa", lapa, "--claude", claude,
                        "--artifacts", artifacts, "--reviewers", *reviewers, "--out", out, expect_ok=False)
            check(again.returncode != 0 and "already exists" in again.stderr, "refuses to overwrite an existing key")
            out2 = tmp / "review_again"
            run("make_review_split.py", "--gemma", gemma, "--lapa", lapa, "--claude", claude,
                "--artifacts", artifacts, "--reviewers", *reviewers, "--out", out2)
            check(pd.read_csv(out / "key/review_key.csv").equals(pd.read_csv(out2 / "key/review_key.csv")),
                  "same seed gives an identical key")
            scenario("clear_winner", out, key, data, reviewers, {"gemma": -0.5, "lapa": 0.0, "claude": 0.5}, "winner", tmp)
            scenario("equal_models", out, key, data, reviewers, {"gemma": 0.0, "lapa": 0.0, "claude": 0.0}, "tie", tmp)
            scenario("consistent_reviewers", out, key, data, reviewers, {"gemma": -0.5, "lapa": 0.0, "claude": 0.5},
                     "winner", tmp, biases=(0.0, 0.0, 0.0), noise=0.1, expect_level="reliable")
            scenario("inconsistent_reviewers", out, key, data, reviewers, {"gemma": -0.5, "lapa": 0.0, "claude": 0.5},
                     "recalibrate", tmp, biases=(0.0, 1.0, -1.0), noise=2.0, expect_level="insufficient")
            broken_returns(out, key, data, reviewers, tmp)
        else:
            out = Path(args.real).resolve()
            manifest = json.loads((out / "key/split_manifest.json").read_text())
            inputs = manifest["inputs"]
            print(f"REAL SPLIT in {out}")
            paths = {}
            for name, info in inputs.items():
                path = Path(info["path"]) if Path(info["path"]).is_absolute() else ROOT / info["path"]
                if name == "claude" and not path.exists():
                    path = resolve_claude_csv(ROOT, None)  # e.g. split made on another machine: use HF cache copy
                paths[name] = path
                check(path.exists() and sha256(path) == info["sha256"], f"{name} input unchanged since the split ({path.name})")
            data = load_translations(paths["gemma"], paths["lapa"], paths["claude"], None)
            reviewers = manifest["reviewers"]
            key = check_split(out, data, reviewers, manifest["shared_prompts"], manifest["prompts_per_reviewer"])
            print("Dry run of scoring on SIMULATED answers (real files are not modified)")
            scenario("real_files_simulated", out, key, data, reviewers, {"gemma": -0.3, "lapa": 0.0, "claude": 0.3}, "winner", tmp)
    finally:
        if args.keep:
            print(f"Kept {tmp}")
        else:
            shutil.rmtree(tmp, ignore_errors=True)
    print(f"\n{'ALL CHECKS PASSED' if not FAILURES else f'{len(FAILURES)} CHECK(S) FAILED'}")
    for f in FAILURES:
        print("  -", f)
    return 1 if FAILURES else 0


if __name__ == "__main__":
    sys.exit(main())
