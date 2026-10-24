#!/usr/bin/env python3
"""Create blinded review packages for the Gemma / Lapa / Claude Ukrainian translations.

Each reviewer gets one Excel file with `--per-reviewer` prompts (default 200). Every prompt shows the English
text and its three translations as options A/B/C in a balanced random order; reviewers never see which model
produced which option. A shared block (default 90 prompts) is given to every reviewer to measure agreement; the
rest of each reviewer's prompts are their own. Records whose English text is identical are kept together, and
each block keeps roughly the JBB/AdvBench mix of the whole set.

Outputs (default under data/safety/review/):
  for_reviewers/review_<name>.xlsx   send these, one per reviewer
  key/review_key.csv                 SECRET: option -> record_id + model; never send to reviewers
  key/split_manifest.json            SECRET: seed, counts, input checksums, records not reviewed

Usage:
  python src/translation_review/make_review_split.py --reviewers anna bohdan ivan
"""

from __future__ import annotations

import argparse
import json
import random
import string
import sys
from collections import Counter, deque
from datetime import datetime, timezone
from itertools import permutations
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).resolve().parent))
from common import (  # noqa: E402
    ISSUES,
    OPTIONS,
    RANK_VALUES,
    REVIEW_SHEET,
    REVIEWER_COLUMNS,
    ROOT,
    SCORE_VALUES,
    SYSTEMS,
    load_translations,
    normalise_goal,
    resolve_claude_csv,
    sha256,
)

ID_ALPHABET = "".join(c for c in string.ascii_uppercase + string.digits if c not in "O0I1L")


def parse_args(argv=None):
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--gemma", default=str(ROOT / "data/safety/translations/prompt-pairs-uk-gemma.csv"))
    p.add_argument("--lapa", default=str(ROOT / "data/safety/translations/prompt-pairs-uk-lapa.csv"))
    p.add_argument("--claude", default=None, help="pinned prompt-pairs.csv; default: local copy or HF cache")
    p.add_argument("--artifacts", default=str(ROOT / "configs/artifacts.json"),
                   help="checksum source for the Claude file; 'none' skips the check")
    p.add_argument("--reviewers", nargs="+", default=["reviewer_1", "reviewer_2", "reviewer_3"])
    p.add_argument("--shared", type=int, default=90, help="prompts every reviewer sees")
    p.add_argument("--per-reviewer", type=int, default=200, help="prompts per reviewer, shared included")
    p.add_argument("--seed", type=int, default=20261011)
    p.add_argument("--out", default=str(ROOT / "data/safety/review"))
    p.add_argument("--force", action="store_true", help="overwrite an existing key (only before files were sent)")
    p.add_argument("--offline", action="store_true", help="do not contact Hugging Face for the Claude file")
    return p.parse_args(argv)


def build_units(data: pd.DataFrame) -> list[dict]:
    """Group records with identical English text; a unit is always assigned as a whole."""
    groups: dict[str, list[int]] = {}
    for rid, goal in data.goal.items():
        groups.setdefault(normalise_goal(goal), []).append(int(rid))
    units = []
    for rids in groups.values():
        sources = set(data.loc[rids, "source"])
        units.append({"records": sorted(rids), "stratum": sources.pop() if len(sources) == 1 else "mixed"})
    return units


def interleave(units: list[dict], rng: random.Random) -> list[dict]:
    """Shuffle within each stratum, then spread strata evenly so any prefix keeps the overall mix."""
    keyed = []
    for stratum in sorted({u["stratum"] for u in units}):
        members = [u for u in units if u["stratum"] == stratum]
        rng.shuffle(members)
        n = len(members)
        keyed += [((i + rng.random()) / n, u) for i, u in enumerate(members)]
    return [u for _, u in sorted(keyed, key=lambda pair: pair[0])]


def fill_blocks(units: list[dict], sizes: list[tuple[str, int]]):
    """Take units in order until each block holds exactly its size (in records)."""
    pending = deque(units)
    blocks = {}
    for name, size in sizes:
        chosen, skipped, total = [], [], 0
        while pending and total < size:
            unit = pending.popleft()
            if total + len(unit["records"]) <= size:
                chosen.append(unit)
                total += len(unit["records"])
            else:
                skipped.append(unit)
        if total != size:
            raise SystemExit(f"Could not fill block {name} to exactly {size} records (got {total})")
        pending = deque(skipped + list(pending))
        blocks[name] = chosen
    return blocks, list(pending)


def make_ids(rng: random.Random, n: int, prefix: str, length: int) -> list[str]:
    ids = set()
    while len(ids) < n:
        ids.add(prefix + "".join(rng.choice(ID_ALPHABET) for _ in range(length)))
    out = sorted(ids)
    rng.shuffle(out)
    return out


INSTRUCTIONS = [
    ("Translation review", "title"),
    ("", None),
    ("Content warning", "head"),
    ("The English prompts come from AI-safety benchmarks (AdvBench, JailbreakBench) and describe harmful requests: "
     "violence, abuse, hate speech, crime, self-harm and similar. You are only judging the translation. You may skip "
     "any item (leave it empty and write 'skipped' in the comment) or stop at any time; tell the coordinator.", None),
    ("", None),
    ("Your task", "head"),
    ("Each prompt (one prompt_id, three rows) shows an English request and three Ukrainian translations, options A, B "
     "and C, in random order. Rate every option on its own, then rank the three against each other.", None),
    ("A good translation is the same request in natural Ukrainian: same meaning, intent, constraints, numbers and names; "
     "nothing softened, sanitised, intensified, added or dropped; informal singular imperative (Напиши, Створи, Поясни), "
     "not formal/plural (Напишіть).", None),
    ("", None),
    ("Columns to fill (dropdowns)", "head"),
    ("adequacy  4 = same meaning, nothing lost/added | 3 = minor nuance lost | 2 = important part wrong or missing | "
     "1 = different request / not a translation", None),
    ("fluency   4 = natural Ukrainian, right register | 3 = understandable, slightly awkward or formal | "
     "2 = clumsy, calque or several errors | 1 = hard to read", None),
    ("rank      1 = best of the three, 3 = worst. Ties allowed (e.g. 1, 1, 3). Identical options get the same rank "
     "and scores.", None),
    ("issue     the main problem of this option, or 'none'", None),
    ("comment   optional; useful for 'other' or to explain a score of 1-2", None),
    ("", None),
    ("Please", "head"),
    ("Do not edit the English, the translations or the 'code' column, and do not sort or delete rows. "
     "Save as .xlsx with the same file name and send it back to the coordinator. Work independently; do not discuss "
     "items with other reviewers until everyone has finished.", None),
]


def write_reviewer_workbook(path: Path, rows: list[dict], reviewer: str, n_prompts: int):
    wb = Workbook()
    ins = wb.active
    ins.title = "Instructions"
    ins.column_dimensions["A"].width = 130
    for r, (text, kind) in enumerate(INSTRUCTIONS + [("", None), (f"Reviewer: {reviewer} · {n_prompts} prompts · {len(rows)} rows", None)], start=1):
        cell = ins.cell(row=r, column=1, value=text)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if kind == "title":
            cell.font = Font(size=16, bold=True)
        elif kind == "head":
            cell.font = Font(size=12, bold=True)

    ws = wb.create_sheet(REVIEW_SHEET)
    ws.append(list(REVIEWER_COLUMNS))
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9E2EC")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    shade = PatternFill("solid", fgColor="F3F5F7")
    for row in rows:
        ws.append([row[c] for c in REVIEWER_COLUMNS])
        if row["item_no"] % 2 == 0:
            for cell in ws[ws.max_row]:
                cell.fill = shade
    widths = {"item_no": 9, "prompt_id": 11, "option": 7, "english_prompt": 48, "ukrainian_translation": 58,
              "adequacy": 9, "fluency": 9, "rank": 7, "issue": 20, "comment": 30, "code": 10}
    col_index = {name: i + 1 for i, name in enumerate(REVIEWER_COLUMNS)}
    for name, width in widths.items():
        ws.column_dimensions[ws.cell(row=1, column=col_index[name]).column_letter].width = width
    for row_cells in ws.iter_rows(min_row=2):
        for cell in row_cells:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "D2"
    last = ws.max_row

    def letter(name):
        return ws.cell(row=1, column=col_index[name]).column_letter

    for name, values in (("adequacy", SCORE_VALUES), ("fluency", SCORE_VALUES), ("rank", RANK_VALUES), ("issue", ISSUES)):
        dv = DataValidation(type="list", formula1='"' + ",".join(map(str, values)) + '"', allow_blank=True,
                            showErrorMessage=True, errorTitle="Invalid value",
                            error="Choose a value from the dropdown list.")
        ws.add_data_validation(dv)
        dv.add(f"{letter(name)}2:{letter(name)}{last}")
    wb.active = 1
    wb.save(path)


def portable(path) -> str:
    """Path relative to the project root when inside it, so the manifest works on another machine."""
    try:
        return str(Path(path).resolve().relative_to(ROOT))
    except ValueError:
        return str(path)


def main(argv=None):
    args = parse_args(argv)
    out = Path(args.out)
    key_dir, send_dir = out / "key", out / "for_reviewers"
    key_path = key_dir / "review_key.csv"
    if key_path.exists() and not args.force:
        raise SystemExit(f"{key_path} already exists. Regenerating would invalidate files already sent to reviewers; "
                         "use --force only if no file has been sent yet.")
    if len(set(args.reviewers)) != len(args.reviewers):
        raise SystemExit("Reviewer names must be unique")
    own = args.per_reviewer - args.shared
    if own <= 0:
        raise SystemExit("--per-reviewer must be larger than --shared")

    claude_path = resolve_claude_csv(ROOT, args.claude, offline=args.offline)
    artifacts = None if args.artifacts.lower() == "none" else Path(args.artifacts)
    data = load_translations(Path(args.gemma), Path(args.lapa), claude_path, artifacts)
    needed = args.shared + own * len(args.reviewers)
    if needed > len(data):
        raise SystemExit(f"{needed} records needed but only {len(data)} available")

    rng = random.Random(args.seed)
    units = interleave(build_units(data), rng)
    sizes = [("shared", args.shared)] + [(f"own:{r}", own) for r in args.reviewers]
    blocks, leftover = fill_blocks(units, sizes)
    record_block = {rid: name for name, us in blocks.items() for u in us for rid in u["records"]}

    prompt_ids = dict(zip(data.index, make_ids(rng, len(data), "P-", 4)))
    codes = dict(zip([(int(rid), s) for rid in data.index for s in SYSTEMS], make_ids(rng, 3 * len(data), "T-", 6)))
    perms = list(permutations(SYSTEMS))

    key_rows, position_counts = [], {}
    send_dir.mkdir(parents=True, exist_ok=True)
    key_dir.mkdir(parents=True, exist_ok=True)
    for reviewer in args.reviewers:
        records = [rid for u in blocks["shared"] + blocks[f"own:{reviewer}"] for rid in u["records"]]
        rng.shuffle(records)
        orders = [perms[i % len(perms)] for i in range(len(records))]
        rng.shuffle(orders)
        rows = []
        for item_no, (rid, order) in enumerate(zip(records, orders), start=1):
            rec = data.loc[rid]
            texts = {s: rec[f"uk_{s}"] for s in SYSTEMS}
            for option, system in zip(OPTIONS, order):
                same = [o for o, s in zip(OPTIONS, order) if s != system and texts[s] == texts[system]]
                rows.append({"item_no": item_no, "prompt_id": prompt_ids[rid], "option": option,
                             "english_prompt": rec.goal, "ukrainian_translation": texts[system],
                             "adequacy": None, "fluency": None, "rank": None, "issue": None, "comment": None,
                             "code": codes[(int(rid), system)]})
                key_rows.append({"reviewer": reviewer, "item_no": item_no, "prompt_id": prompt_ids[rid],
                                 "option": option, "code": codes[(int(rid), system)], "record_id": int(rid),
                                 "source": rec.source, "system": system,
                                 "block": "shared" if record_block[rid] == "shared" else "own",
                                 "identical_to_options": "".join(same)})
                position_counts[(reviewer, system, option)] = position_counts.get((reviewer, system, option), 0) + 1
        write_reviewer_workbook(send_dir / f"review_{reviewer}.xlsx", rows, reviewer, len(records))
        print(f"{reviewer}: {len(records)} prompts, {len(rows)} rows -> {send_dir / f'review_{reviewer}.xlsx'}")

    key = pd.DataFrame(key_rows)
    key.to_csv(key_path, index=False, encoding="utf-8")
    not_reviewed = sorted(rid for u in leftover for rid in u["records"])
    jbb_share = lambda rids: round(float((data.loc[rids, "source"] == "JBB").mean()), 3) if rids else None
    manifest = {
        "created_utc": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "seed": args.seed,
        "reviewers": args.reviewers,
        "shared_prompts": args.shared,
        "prompts_per_reviewer": args.per_reviewer,
        "inputs": {
            "gemma": {"path": portable(args.gemma), "sha256": sha256(Path(args.gemma))},
            "lapa": {"path": portable(args.lapa), "sha256": sha256(Path(args.lapa))},
            "claude": {"path": portable(claude_path), "sha256": sha256(claude_path),
                       "note": "pinned prompt-pairs.csv (configs/artifacts.json)"},
        },
        "records_total": int(len(data)),
        "records_reviewed": int(len(record_block)),
        "records_not_reviewed": not_reviewed,
        "blocks": {name: {"records": sorted(r for u in us for r in u["records"]),
                          "jbb_share": jbb_share([r for u in us for r in u["records"]])}
                   for name, us in blocks.items()},
        "jbb_share_overall": jbb_share(list(data.index)),
        "duplicate_english_groups": [u["records"] for u in units if len(u["records"]) > 1],
        "option_position_counts": {f"{r}/{s}/{o}": n for (r, s, o), n in sorted(position_counts.items())},
        "scales": {"adequacy": list(SCORE_VALUES), "fluency": list(SCORE_VALUES), "rank": list(RANK_VALUES),
                   "issue": list(ISSUES)},
        "note": "SECRET. key/ unblinds the review; never send it to reviewers.",
    }
    (key_dir / "split_manifest.json").write_text(json.dumps(manifest, indent=2, ensure_ascii=False), encoding="utf-8")
    print(f"Key: {key_path}  (keep secret)")
    print(f"Records reviewed: {len(record_block)} of {len(data)}; shared {args.shared}, own {own} per reviewer; "
          f"not reviewed: {len(not_reviewed)}")
    print("JBB share - overall:", manifest["jbb_share_overall"],
          "| blocks:", {n: b["jbb_share"] for n, b in manifest["blocks"].items()})
    return manifest


if __name__ == "__main__":
    main()
